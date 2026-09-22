"""Build the Amsterdam UMC-branded DeepPeak webinar as a native Keynote file."""

from pathlib import Path
import subprocess

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).parent
PPTX_PATH = ROOT / "presentation.pptx"
KEYNOTE_PATH = ROOT / "presentation.key"
REVIEW_PDF_PATH = Path("/private/tmp/deeppeak-webinar-review.pdf")
LOGO_PATH = ROOT / "amsterdam-umc-logo-presentation.png"
MOTIF_PATH = ROOT / "pulse-overlap-motif.png"
CNN_TRACE_PATH = ROOT / "cyto-cnn-trace.png"
ARRIVAL_TIME_PATH = ROOT / "cyto-arrival-time.png"
AMPLITUDES_PATH = ROOT / "cyto-amplitudes.png"
THROUGHPUT_PATH = ROOT / "cyto-throughput.png"
PRESENTER_NAME = "Martin Poinsinet de Sivry"

TEAL = "006D77"
DARK_TEAL = "004854"
CORAL = "EF8354"
INK = "102A33"
MUTED = "56727A"
MINT = "F2F8F7"
PALE_TEAL = "DCEDEA"
PAPER = "F8FBFA"
SEAFOAM = "CFE5E1"
SUN = "F2B84B"
WHITE = "FFFFFF"

SECTION_STYLES = {
    "package": {
        "label": "01  DEEPPEAK",
        "accent": TEAL,
        "canvas": "F1F8F7",
        "rail": DARK_TEAL,
        "soft": PALE_TEAL,
    },
    "problem": {
        "label": "02  THE PROBLEM",
        "accent": CORAL,
        "canvas": "FCF6F2",
        "rail": "7F3F31",
        "soft": "FBE5DC",
    },
    "method": {
        "label": "03  THE METHOD",
        "accent": TEAL,
        "canvas": "F4F9F8",
        "rail": "245B62",
        "soft": "DDEDEA",
    },
    "evidence": {
        "label": "04  THE EVIDENCE",
        "accent": "B56B12",
        "canvas": "FBF8F1",
        "rail": INK,
        "soft": "F4E8CC",
    },
    "action": {
        "label": "05  APPLY IT",
        "accent": CORAL,
        "canvas": "F5F8F8",
        "rail": DARK_TEAL,
        "soft": "E3EFED",
    },
}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def rect(slide, x, y, width, height, color, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, color, width=1.2):
    shape = slide.shapes.add_connector(
        1, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    return shape


def image(slide, path, x, y, width, height=None):
    if height is None:
        return slide.shapes.add_picture(
            str(path), Inches(x), Inches(y), width=Inches(width)
        )
    return slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), Inches(width), Inches(height)
    )


def text(
    slide, value, x, y, width, height, size, color=INK, bold=False, align=PP_ALIGN.LEFT
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = align
    paragraph.font.name = "Avenir Next"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = rgb(color)
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    return box


def code(slide, value, x, y, width, height, size=12, color="D8EEEE"):
    box = text(slide, value, x, y, width, height, size, color)
    for paragraph in box.text_frame.paragraphs:
        paragraph.font.name = "Menlo"
        for run in paragraph.runs:
            run.font.name = "Menlo"
    return box


def add_logo(slide, x=10.78, y=0.33, width=1.95):
    slide.shapes.add_picture(str(LOGO_PATH), Inches(x), Inches(y), width=Inches(width))


def add_footer(slide, number, dark=False, section=None):
    color = "BFD0D2" if dark else MUTED
    rule_color = "2A6670" if dark else SEAFOAM
    rect(slide, 0.82, 6.70, 11.70, 0.018, rule_color)
    text(slide, PRESENTER_NAME, 0.82, 6.94, 3.10, 0.18, 9, color, bold=True)
    footer_label = (
        SECTION_STYLES[section]["label"] if section else "Amsterdam UMC | DeepPeak"
    )
    text(slide, footer_label, 4.10, 6.94, 3.0, 0.18, 9, color, bold=bool(section))
    text(
        slide, f"{number:02d}", 12.05, 6.94, 0.35, 0.18, 9, color, align=PP_ALIGN.RIGHT
    )


def add_section_layout(slide, section, eyebrow, title, number):
    """Apply a named, reusable section layout to a content slide."""
    style = SECTION_STYLES[section]
    rect(slide, 0, 0, 13.333, 7.5, style["canvas"])
    rect(slide, 0, 0, 0.18, 7.5, style["rail"])
    rect(slide, 0.18, 0, 13.153, 0.12, style["accent"])
    rect(slide, 0.82, 0.42, 1.62, 0.34, style["accent"], radius=True)
    text(
        slide,
        style["label"],
        0.82,
        0.51,
        1.62,
        0.14,
        8,
        WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    text(slide, eyebrow, 2.72, 0.50, 6.6, 0.20, 9, style["accent"], bold=True)
    text(slide, title, 0.82, 0.98, 10.25, 0.60, 24, INK, bold=True)
    rect(slide, 11.18, 0.36, 1.35, 0.48, WHITE, radius=True)
    add_logo(slide, 11.29, 0.43, 1.12)
    add_footer(slide, number, section=section)


def add_slide_heading(slide, eyebrow, title, number, section="method"):
    add_section_layout(slide, section, eyebrow, title, number)


def add_card(
    slide, x, y, width, height, eyebrow, heading, body, accent=TEAL, dark=False
):
    """Add a reusable subject card with a strong semantic hierarchy."""
    background = DARK_TEAL if dark else WHITE
    heading_color = WHITE if dark else INK
    body_color = "CFE0E1" if dark else MUTED
    rect(slide, x, y, width, height, background, radius=True)
    rect(slide, x, y, 0.09, height, accent, radius=True)
    rect(slide, x + 0.28, y + 0.28, 0.42, 0.07, accent)
    text(slide, eyebrow, x + 0.28, y + 0.54, width - 0.56, 0.18, 9, accent, bold=True)
    text(
        slide,
        heading,
        x + 0.28,
        y + 0.94,
        width - 0.56,
        0.50,
        18,
        heading_color,
        bold=True,
    )
    text(slide, body, x + 0.28, y + 1.58, width - 0.56, height - 1.82, 12, body_color)


def add_code_card(slide, value, x, y, width, height, caption):
    """Add the dedicated code-only visual component."""
    rect(slide, x, y, width, height, "0B2932", radius=True)
    rect(slide, x, y, width, 0.52, "123B45", radius=True)
    for index, color in enumerate((CORAL, SUN, "65B9AD")):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x + 0.28 + index * 0.25),
            Inches(y + 0.18),
            Inches(0.11),
            Inches(0.11),
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb(color)
        dot.line.fill.background()
    text(slide, caption, x + 1.12, y + 0.17, width - 1.42, 0.15, 8, "A8CCCB", bold=True)
    code(slide, value, x + 0.34, y + 0.78, width - 0.68, height - 1.02, 10, "D8EEEE")


def add_stage(slide, x, heading, body, accent):
    add_card(
        slide,
        x,
        3.02,
        3.45,
        2.42,
        "MEASUREMENT REGIME",
        heading,
        body,
        accent,
    )


def add_figure_card(slide, path, x, y, width, height, label, accent=TEAL):
    rect(slide, x, y, width, height, WHITE, radius=True)
    rect(slide, x, y, width, 0.42, accent, radius=True)
    text(slide, label, x + 0.24, y + 0.13, width - 0.48, 0.14, 8, WHITE, bold=True)
    image(slide, path, x + 0.18, y + 0.58, width - 0.36, height - 0.76)


def build_presentation() -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    # 1. Title
    slide = presentation.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, PAPER)
    rect(slide, 0, 0, 0.17, 7.5, TEAL)
    rect(slide, 7.48, 0, 5.853, 7.5, DARK_TEAL)
    rect(slide, 0.84, 1.10, 0.13, 4.20, CORAL)
    rect(slide, 9.76, 0.56, 2.54, 0.67, WHITE, radius=True)
    add_logo(slide, 9.92, 0.67, 2.20)
    text(slide, "DEEPPEAK WEBINAR", 1.22, 1.20, 4.8, 0.28, 11, TEAL, bold=True)
    text(
        slide,
        "High-Throughput\nExtracellular Vesicle Detection",
        1.22,
        1.75,
        5.95,
        1.34,
        31,
        INK,
        bold=True,
    )
    text(
        slide,
        "The open-source package for overlapping-pulse recovery",
        1.22,
        3.42,
        5.55,
        0.34,
        16,
        CORAL,
        bold=True,
    )
    text(
        slide,
        "From realistic simulation and neural detection to validated EV measurements",
        1.22,
        4.16,
        5.20,
        0.58,
        15,
        MUTED,
    )
    rect(slide, 7.94, 2.12, 4.86, 2.95, WHITE, radius=True)
    slide.shapes.add_picture(
        str(MOTIF_PATH), Inches(8.13), Inches(2.39), width=Inches(4.46)
    )
    text(
        slide,
        "SIGNAL OVERLAP | EVENT RECOVERY",
        8.13,
        5.39,
        4.46,
        0.2,
        10,
        "89BDB9",
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    text(slide, PRESENTER_NAME, 1.22, 6.77, 3.12, 0.18, 10, MUTED, bold=True)
    text(slide, "Amsterdam UMC | DeepPeak", 4.52, 6.77, 2.70, 0.18, 10, MUTED)
    text(slide, "01", 12.05, 6.77, 0.35, 0.18, 10, "89BDB9", align=PP_ALIGN.RIGHT)

    # 2. Roadmap
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "20-MINUTE WEBINAR",
        "From the DeepPeak package to validated recovery",
        2,
        "package",
    )
    roadmap = [
        (
            "01",
            "Meet DeepPeak",
            "Package architecture and a reusable Python workflow.",
            "00:00-05:00",
        ),
        (
            "02",
            "The bottleneck",
            "Why overlap erases information—and how FLASH responds.",
            "05:00-10:00",
        ),
        (
            "03",
            "The case study",
            "Timing, amplitude, and throughput checks on data.",
            "10:00-17:00",
        ),
        (
            "04",
            "Your workflow",
            "Validation controls and EV translation.",
            "17:00-20:00",
        ),
    ]
    for index, (number, heading, body, timing) in enumerate(roadmap):
        y = 2.15 + (index * 0.95)
        rect(
            slide, 0.82, y, 0.52, 0.52, CORAL if index in (1, 2) else TEAL, radius=True
        )
        text(
            slide,
            number,
            0.82,
            y + 0.15,
            0.52,
            0.15,
            9,
            WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        text(slide, heading, 1.62, y + 0.02, 2.25, 0.24, 17, INK, bold=True)
        text(slide, body, 3.92, y + 0.04, 5.55, 0.28, 13, MUTED)
        text(
            slide,
            timing,
            10.08,
            y + 0.05,
            1.48,
            0.22,
            11,
            TEAL,
            bold=True,
            align=PP_ALIGN.RIGHT,
        )

    # 3. Package overview
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "THE PACKAGE",
        "DeepPeak turns 1D detector traces into measurable events",
        3,
        "package",
    )
    text(
        slide,
        "An open-source Python toolkit for building, testing, and applying peak-recovery workflows.",
        0.82,
        2.02,
        10.9,
        0.36,
        16,
        MUTED,
    )
    capabilities = [
        ("GENERATE", "Realistic pulse trains, overlap, noise, drift, and blanks."),
        (
            "LEARN",
            "DenseNet, WaveNet, and UNet1D models with reusable training controls.",
        ),
        (
            "ANALYZE",
            "Standard and neural trace analysis, distributions, and comparison metrics.",
        ),
        (
            "VALIDATE",
            "Typed pipelines and structured results for reproducible experiments.",
        ),
    ]
    for index, (heading, body) in enumerate(capabilities):
        x = 0.82 + (index * 3.02)
        accent = CORAL if index == 1 else TEAL
        add_card(
            slide,
            x,
            2.70,
            2.58,
            2.72,
            f"CAPABILITY 0{index + 1}",
            heading,
            body,
            accent,
        )
    text(
        slide,
        "pip install DeepPeak   •   github.com/MartinPdeS/DeepPeak",
        0.82,
        5.92,
        8.8,
        0.25,
        13,
        TEAL,
        bold=True,
    )

    # 4. Package workflow
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "A REUSABLE PYTHON WORKFLOW",
        "Configure the acquisition once; compare methods consistently",
        4,
        "package",
    )
    code_value = (
        "from DeepPeak import (\n"
        "    Gaussian, PoissonCount, SignalGenerator\n"
        ")\n\n"
        "generator = SignalGenerator(sequence_length=256)\n"
        "dataset = generator.generate(\n"
        "    n_samples=2048,\n"
        "    kernel=Gaussian(width=(8, 18)),\n"
        "    peak_count=PoissonCount(bounds=(0, 4)),\n"
        "    noise_std=(0.01, 0.12),\n"
        ")"
    )
    add_code_card(
        slide,
        code_value,
        1.18,
        2.10,
        6.38,
        3.98,
        "PYTHON  •  SYNTHETIC TRAINING DATA",
    )
    outcomes = [
        ("CONTROL", "Known event times and amplitudes"),
        ("COMPARE", "Standard versus neural detection"),
        ("REPORT", "Recovery, bias, timing, throughput"),
    ]
    for index, (heading, body) in enumerate(outcomes):
        y = 2.20 + (index * 1.24)
        accent = CORAL if index == 1 else TEAL
        rect(slide, 8.42, y, 3.72, 0.94, MINT, radius=True)
        rect(slide, 8.42, y, 0.09, 0.94, accent, radius=True)
        text(slide, heading, 8.76, y + 0.18, 1.10, 0.18, 10, accent, bold=True)
        text(slide, body, 8.76, y + 0.48, 2.88, 0.22, 13, INK, bold=True)
    text(
        slide,
        "The same data model carries simulation truth into evaluation.",
        8.42,
        6.02,
        3.72,
        0.26,
        12,
        MUTED,
        bold=True,
    )

    # 5. Bottleneck
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "THE MEASUREMENT PROBLEM",
        "When throughput rises, pulse overlap becomes the bottleneck",
        5,
        "problem",
    )
    text(
        slide,
        "Conventional peak detection assumes isolated, well-separated events.",
        0.78,
        2.16,
        9.7,
        0.34,
        17,
        MUTED,
    )
    add_stage(
        slide, 0.78, "LOW RATE", "Separated pulses\ncan be counted directly.", TEAL
    )
    add_stage(
        slide,
        4.94,
        "HIGHER RATE",
        "Partially overlapping pulses\nlose timing and count fidelity.",
        CORAL,
    )
    add_stage(
        slide,
        9.10,
        "SATURATION RISK",
        "Merged responses distort\nmeasured event populations.",
        DARK_TEAL,
    )
    line(slide, 4.3, 4.24, 4.72, 4.24, PALE_TEAL, 2.2)
    line(slide, 8.46, 4.24, 8.88, 4.24, PALE_TEAL, 2.2)

    # 6. FLASH pipeline
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "THE FLASH PIPELINE",
        "Localize events first, then recover amplitudes",
        6,
        "method",
    )
    text(
        slide,
        "FLASH combines a CNN event-localization stage with an analytical amplitude solver.",
        0.78,
        2.17,
        10.2,
        0.5,
        18,
        MUTED,
    )
    rect(slide, 0.78, 3.25, 5.55, 2.05, PALE_TEAL, radius=True)
    rect(slide, 7.00, 3.25, 5.55, 2.05, "FBE8E1", radius=True)
    text(slide, "CNN LOCALIZATION", 1.14, 3.68, 2.8, 0.2, 11, TEAL, bold=True)
    text(
        slide,
        "Estimate arrival times\nfrom the noisy detector trace.",
        1.14,
        4.13,
        4.58,
        0.72,
        21,
        INK,
        bold=True,
    )
    text(slide, "ANALYTICAL RECOVERY", 7.36, 3.68, 3.2, 0.2, 11, CORAL, bold=True)
    text(
        slide,
        "Solve amplitudes from the\nknown response and event times.",
        7.36,
        4.13,
        4.6,
        0.72,
        21,
        INK,
        bold=True,
    )

    # 7. Workflow
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "THE MODELING WORKFLOW",
        "DeepPeak provides a realistic training workflow",
        7,
        "method",
    )
    steps = [
        (
            "01",
            "Generate",
            "Pulse-shape variability, Poisson event counts, blanks, noise, and drift.",
        ),
        (
            "02",
            "Train",
            "WaveNet event-location learning with validation-aware controls.",
        ),
        (
            "03",
            "Evaluate",
            "Recovery, timing error, count bias, confidence, and retained throughput.",
        ),
    ]
    for index, (number, heading, body) in enumerate(steps):
        y = 2.26 + (index * 1.2)
        rect(slide, 0.80, y, 0.58, 0.58, CORAL if index == 1 else TEAL, radius=True)
        text(
            slide,
            number,
            0.80,
            y + 0.16,
            0.58,
            0.18,
            11,
            WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        text(slide, heading, 1.68, y + 0.02, 1.8, 0.28, 19, INK, bold=True)
        text(slide, body, 1.68, y + 0.42, 8.9, 0.3, 14, MUTED)
        if index < 2:
            line(slide, 1.09, y + 0.62, 1.09, y + 1.18, PALE_TEAL, 2.2)
    text(
        slide,
        "SignalGenerator + Gaussian + PoissonCount + WaveNet",
        1.68,
        6.10,
        7.5,
        0.24,
        13,
        TEAL,
        bold=True,
    )

    # 8. Training realism
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "REALISM MATTERS",
        "Design the training distribution around the acquisition",
        8,
        "method",
    )
    columns = [
        (
            "PULSE LIBRARY",
            "Widths, amplitudes, positions, and asymmetric shapes reflect expected EV signals.",
        ),
        (
            "BACKGROUND",
            "Vary noise level, baseline level, and baseline drift across trace batches.",
        ),
        (
            "NEGATIVES",
            "Include blank traces so the detector learns when not to call an event.",
        ),
    ]
    for index, (heading, body) in enumerate(columns):
        x = 0.78 + (index * 4.16)
        accent = TEAL if index != 1 else CORAL
        add_card(
            slide,
            x,
            2.54,
            3.48,
            2.78,
            f"TRAINING INPUT 0{index + 1}",
            heading,
            body,
            accent,
        )

    # 9. Experimental system
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "EXPERIMENTAL SYSTEM",
        "Test the method where coincidence is real",
        9,
        "evidence",
    )
    text(
        slide,
        "The CYTO study used high-flow light-scattering measurements from 300 nm polystyrene beads.",
        0.82,
        2.10,
        10.8,
        0.34,
        16,
        MUTED,
    )
    system_cards = [
        (
            "ACQUISITION",
            "BD FACSCanto II flow cell and laser optics\nwith electronics bypassed.",
        ),
        (
            "DIGITIZATION",
            "Side-scatter PMT signal recorded with\na 125 MHz, 14-bit PicoScope.",
        ),
        (
            "REALITY CHECK",
            "A distorted response with a secondary lobe\ncreates a demanding overlap problem.",
        ),
    ]
    for index, (heading, body) in enumerate(system_cards):
        x = 0.82 + (index * 4.05)
        accent = CORAL if index == 2 else "B56B12"
        add_card(
            slide,
            x,
            3.02,
            3.36,
            2.45,
            f"SYSTEM 0{index + 1}",
            heading,
            body,
            accent,
        )

    # 10. CNN evidence
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "CYTO EXPERIMENT",
        "CNN localization resolves dense event sequences",
        10,
        "evidence",
    )
    text(
        slide,
        "At approximately 94k events/s, the model maps the detector trace to sharply localized event support.",
        0.82,
        2.08,
        10.8,
        0.34,
        16,
        MUTED,
    )
    add_figure_card(
        slide,
        CNN_TRACE_PATH,
        0.82,
        2.72,
        7.25,
        3.34,
        "CYTO 2026 | noisy detector trace and CNN event prediction",
        CORAL,
    )
    rect(slide, 8.48, 2.72, 3.90, 3.34, DARK_TEAL, radius=True)
    text(slide, "WHY THIS MATTERS", 8.84, 3.16, 2.9, 0.20, 11, "74D1D0", bold=True)
    text(
        slide,
        "The neural model supplies event timing.\n\nThat timing makes a deterministic amplitude recovery step possible.",
        8.84,
        3.58,
        2.95,
        1.80,
        17,
        WHITE,
        bold=True,
    )

    # 11. Validation
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "THREE-FOLD VALIDATION",
        "Check timing, amplitude, and throughput independently",
        11,
        "evidence",
    )
    text(
        slide,
        "A high-throughput result is useful only when recovered measurements agree with a credible reference.",
        0.78,
        2.18,
        10.8,
        0.4,
        16,
        MUTED,
    )
    metrics = [
        ("RECOVERY", "Precision\nRecall"),
        ("COUNTING", "Count bias\nDistribution shift"),
        ("LOCALIZATION", "Timing error\nSeparation limit"),
        ("USABILITY", "Retained\nthroughput"),
    ]
    for index, (heading, body) in enumerate(metrics):
        x = 0.87 + (index * 3.14)
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(3.32), Inches(2.12), Inches(2.12)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = rgb(PALE_TEAL if index % 2 == 0 else "FBE8E1")
        circle.line.color.rgb = rgb(TEAL if index % 2 == 0 else CORAL)
        circle.line.width = Pt(1.8)
        text(
            slide,
            heading,
            x + 0.30,
            3.83,
            1.52,
            0.2,
            10,
            TEAL if index % 2 == 0 else CORAL,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        text(
            slide,
            body,
            x + 0.22,
            4.28,
            1.68,
            0.55,
            16,
            INK,
            bold=True,
            align=PP_ALIGN.CENTER,
        )

    # 12. Timing and amplitude evidence
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "CYTO RESULTS",
        "Preserve physical timing and reduce amplitude bias",
        12,
        "evidence",
    )
    add_figure_card(
        slide,
        ARRIVAL_TIME_PATH,
        0.82,
        2.22,
        5.45,
        3.94,
        "Arrival-time statistics | low-rate reference",
        TEAL,
    )
    add_figure_card(
        slide,
        AMPLITUDES_PATH,
        6.65,
        2.22,
        5.45,
        3.94,
        "Amplitude distribution | low-rate reference",
        CORAL,
    )
    text(
        slide,
        "At high event rate, the CYTO study found that FLASH preserved plausible arrival-time behavior while reducing coincidence-induced amplitude bias.",
        0.82,
        6.28,
        10.9,
        0.24,
        13,
        MUTED,
        bold=True,
    )

    # 13. Throughput evidence
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "CYTO RESULTS",
        "Extend the validated event-rate range by >10×",
        13,
        "evidence",
    )
    add_figure_card(
        slide,
        THROUGHPUT_PATH,
        0.82,
        2.10,
        8.15,
        4.22,
        "Measured particle flow versus expected throughput",
        CORAL,
    )
    rect(slide, 9.43, 2.10, 2.90, 4.22, DARK_TEAL, radius=True)
    text(
        slide,
        ">10x",
        9.78,
        2.83,
        2.2,
        0.55,
        35,
        "74D1D0",
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    text(
        slide,
        "usable event-rate\nrange in the CYTO\nexperiment",
        9.78,
        3.72,
        2.2,
        0.82,
        17,
        WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    text(
        slide,
        "Interpret this as a validated operating-range extension, not unlimited deconvolution.",
        9.78,
        5.22,
        2.2,
        0.48,
        11,
        "BFD0D2",
        align=PP_ALIGN.CENTER,
    )

    # 14. Quality control
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "QUALITY CONTROL",
        "Operational controls keep neural estimates scientifically useful",
        14,
        "action",
    )
    controls = [
        "Calibrate confidence thresholds against known mixtures or lower-throughput references.",
        "Track pulse-shape, noise, and baseline distributions for acquisition-to-training drift.",
        "Preserve flagged traces for review instead of forcing a count in ambiguous regions.",
        "Report the excluded region and the retained-throughput gain together.",
    ]
    for index, control in enumerate(controls):
        y = 2.22 + (index * 0.94)
        rect(slide, 0.80, y, 0.38, 0.38, TEAL if index % 2 == 0 else CORAL, radius=True)
        text(
            slide,
            f"0{index + 1}",
            0.80,
            y + 0.10,
            0.38,
            0.14,
            8,
            WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        text(slide, control, 1.46, y + 0.02, 9.98, 0.34, 16, INK, bold=True)

    # 15. EV translation
    slide = presentation.slides.add_slide(blank)
    add_slide_heading(
        slide,
        "TRANSLATING TO EXTRACELLULAR VESICLES",
        "Use DeepPeak to move the throughput boundary with evidence",
        15,
        "action",
    )
    text(
        slide,
        "The CYTO result motivates an EV workflow, but every acquisition must establish its own validated operating range.",
        0.82,
        2.10,
        10.9,
        0.35,
        16,
        MUTED,
    )
    translation = [
        (
            "1",
            "Characterize",
            "Acquire low-rate EV traces and extract representative instrument responses.",
        ),
        (
            "2",
            "Train",
            "Generate realistic overlap, noise, drift, and blank examples with DeepPeak.",
        ),
        (
            "3",
            "Validate",
            "Compare timing, amplitude distributions, and counts against a lower-rate reference.",
        ),
    ]
    for index, (number, heading, body) in enumerate(translation):
        x = 0.82 + (index * 4.05)
        rect(slide, x, 3.20, 3.36, 2.38, MINT, radius=True)
        rect(
            slide,
            x + 0.30,
            3.54,
            0.48,
            0.48,
            CORAL if index == 1 else TEAL,
            radius=True,
        )
        text(
            slide,
            number,
            x + 0.30,
            3.68,
            0.48,
            0.15,
            10,
            WHITE,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        text(slide, heading, x + 1.04, 3.57, 1.98, 0.24, 18, INK, bold=True)
        text(slide, body, x + 0.30, 4.34, 2.72, 0.75, 14, MUTED)

    # 16. Close
    slide = presentation.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, DARK_TEAL)
    rect(slide, 0, 0, 0.17, 7.5, CORAL)
    rect(slide, 9.73, 0.56, 2.54, 0.67, WHITE, radius=True)
    add_logo(slide, 9.89, 0.67, 2.20)
    text(slide, "TAKEAWAY", 0.90, 1.12, 4.2, 0.26, 11, "74D1D0", bold=True)
    text(
        slide,
        "Increase usable throughput\nby recovering resolvable overlap.",
        0.90,
        1.70,
        8.2,
        1.26,
        31,
        WHITE,
        bold=True,
    )
    rect(slide, 0.90, 3.48, 1.15, 0.05, SUN)
    text(
        slide,
        "DeepPeak supports a controlled workflow: realistic generation, neural event localization, and validation against a reference.",
        0.90,
        3.96,
        8.55,
        0.75,
        18,
        "D2E7E5",
    )
    text(
        slide,
        "Next: add acquisition-specific traces, reference data, and validated metrics.",
        0.90,
        5.45,
        8.7,
        0.32,
        15,
        "74D1D0",
        bold=True,
    )
    text(
        slide,
        "github.com/MartinPdeS/DeepPeak",
        9.74,
        5.45,
        2.58,
        0.26,
        12,
        "74D1D0",
        bold=True,
        align=PP_ALIGN.RIGHT,
    )
    add_footer(slide, 16, dark=True)

    presentation.save(PPTX_PATH)


def convert_to_keynote() -> None:
    if KEYNOTE_PATH.exists():
        KEYNOTE_PATH.unlink()
    if REVIEW_PDF_PATH.exists():
        REVIEW_PDF_PATH.unlink()
    subprocess.run(["open", "-a", "Keynote"], check=True)
    script = """
on run argv
    set sourceFile to POSIX file (item 1 of argv)
    set outputFile to POSIX file (item 2 of argv)
    set reviewFile to POSIX file (item 3 of argv)
    tell application "Keynote"
        activate
        set importedDocument to open sourceFile
        save importedDocument in outputFile
        export importedDocument to reviewFile as PDF
        close importedDocument saving no
    end tell
end run
"""
    subprocess.run(
        [
            "osascript",
            "-e",
            script,
            str(PPTX_PATH),
            str(KEYNOTE_PATH),
            str(REVIEW_PDF_PATH),
        ],
        check=True,
    )
    PPTX_PATH.unlink()


if __name__ == "__main__":
    build_presentation()
    convert_to_keynote()
